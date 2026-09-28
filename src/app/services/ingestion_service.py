import csv
import json
import io
import uuid
import hashlib
from datetime import datetime
from typing import Dict, Any, List
from openpyxl import load_workbook
from pypdf import PdfReader

from app.models.schemas import Entity, Relationship, EntityType, RelationshipType, Evidence
from app.db.graph_db import db
from app.services.extraction_service import EntityExtractor

class IngestionService:
    @staticmethod
    def process_file(content_bytes: bytes, filename: str, case_id: str) -> Dict[str, Any]:
        """Unified parser supporting CSV, XLSX, PDF, JSON, TXT, LOG files."""
        fn_lower = filename.lower()

        if fn_lower.endswith(".csv"):
            content_str = content_bytes.decode("utf-8", errors="ignore")
            return IngestionService.process_csv(content_str, filename, case_id)
        elif fn_lower.endswith(".xlsx") or fn_lower.endswith(".xls"):
            return IngestionService.process_excel(content_bytes, filename, case_id)
        elif fn_lower.endswith(".pdf"):
            return IngestionService.process_pdf(content_bytes, filename, case_id)
        elif fn_lower.endswith(".json"):
            content_str = content_bytes.decode("utf-8", errors="ignore")
            return IngestionService.process_json(content_str, filename, case_id)
        else:
            # Plain text / logs / notes
            content_str = content_bytes.decode("utf-8", errors="ignore")
            return IngestionService.process_text(content_str, filename, case_id)

    @staticmethod
    def process_csv(content: str, filename: str, case_id: str) -> Dict[str, Any]:
        reader = csv.DictReader(io.StringIO(content))
        processed_count = 0
        entities_created = 0
        relationships_created = 0

        for row in reader:
            processed_count += 1
            rec_id = row.get("transaction_id") or row.get("call_id") or row.get("id") or f"REC-{processed_count}"

            row_str = " ".join([f"{k}:{v}" for k, v in row.items() if v])
            e_list, r_list, _ = EntityExtractor.extract_from_text(row_str, case_id, filename, rec_id)
            entities_created += len(e_list)
            relationships_created += len(r_list)

            src_acc = row.get("source_account") or row.get("from_account") or row.get("caller")
            dst_acc = row.get("destination_account") or row.get("to_account") or row.get("receiver")

            if src_acc and dst_acc:
                e_src_id = f"ENT-ACC-{src_acc}" if "ACC" in src_acc else f"ENT-PHONE-{src_acc.replace('+', '')}"
                e_dst_id = f"ENT-ACC-{dst_acc}" if "ACC" in dst_acc else f"ENT-PHONE-{dst_acc.replace('+', '')}"

                rel_type = RelationshipType.TRANSFERRED_TO if "ACC" in src_acc else RelationshipType.CALLED
                rel = Relationship(
                    relationship_id=f"REL-{rec_id}",
                    source_entity=e_src_id,
                    relationship_type=rel_type,
                    target_entity=e_dst_id,
                    case_id=case_id,
                    source_record_id=rec_id,
                    timestamp=row.get("timestamp") or row.get("date") or datetime.now().isoformat(),
                    confidence=1.0,
                    extraction_method="structured_csv"
                )
                db.add_relationship(rel)
                relationships_created += 1

        return {
            "records_processed": processed_count,
            "entities_discovered": entities_created,
            "relationships_discovered": relationships_created,
            "status": "COMPLETED"
        }

    @staticmethod
    def process_excel(content_bytes: bytes, filename: str, case_id: str) -> Dict[str, Any]:
        """Parses Excel bank statements & transaction sheets using openpyxl."""
        excel_file = io.BytesIO(content_bytes)
        wb = load_workbook(excel_file, data_only=True)
        sheet = wb.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return {"records_processed": 0, "entities_discovered": 0, "relationships_discovered": 0, "status": "EMPTY"}

        headers = [str(cell) if cell is not None else "" for cell in rows[0]]
        full_text_lines = []
        for r in rows[1:]:
            line = " ".join([f"{headers[i]}:{r[i]}" for i in range(min(len(headers), len(r))) if r[i] is not None])
            full_text_lines.append(line)

        combined_text = "\n".join(full_text_lines)
        return IngestionService.process_text(combined_text, filename, case_id)

    @staticmethod
    def process_pdf(content_bytes: bytes, filename: str, case_id: str) -> Dict[str, Any]:
        """Extracts text from PDF field reports & officer notes using pypdf."""
        pdf_file = io.BytesIO(content_bytes)
        reader = PdfReader(pdf_file)
        full_text = ""
        for page in reader.pages:
            full_text += page.extract_text() + "\n"

        return IngestionService.process_text(full_text, filename, case_id)

    @staticmethod
    def process_json(content: str, filename: str, case_id: str) -> Dict[str, Any]:
        data = json.loads(content)
        records = data if isinstance(data, list) else [data]
        entities_created = 0
        relationships_created = 0

        for idx, rec in enumerate(records):
            rec_id = rec.get("call_id") or rec.get("id") or f"REC-JSON-{idx+1}"
            rec_str = json.dumps(rec)
            e_list, r_list, _ = EntityExtractor.extract_from_text(rec_str, case_id, filename, rec_id)
            entities_created += len(e_list)
            relationships_created += len(r_list)

        return {
            "records_processed": len(records),
            "entities_discovered": entities_created,
            "relationships_discovered": relationships_created,
            "status": "COMPLETED"
        }

    @staticmethod
    def process_text(content: str, filename: str, case_id: str) -> Dict[str, Any]:
        entities, relationships, evidence = EntityExtractor.extract_from_text(
            text=content,
            case_id=case_id,
            source_file=filename,
            source_record_id=f"TXT-{hashlib.md5(content.encode()).hexdigest()[:8]}"
        )
        return {
            "records_processed": 1,
            "entities_discovered": len(entities),
            "relationships_discovered": len(relationships),
            "evidence_id": evidence.evidence_id,
            "status": "COMPLETED"
        }
